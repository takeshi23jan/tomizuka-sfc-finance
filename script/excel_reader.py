#!/usr/bin/env python3
"""Dump non-empty cell contents of receipt-related Excel files in hamakuru for review."""
from __future__ import annotations

from pathlib import Path

import openpyxl

HAMAKURU_DIR = Path(__file__).parent.parent / "hamakuru"

TARGET_FILES = [
    "08ryoushushodaishi.xlsx",
    "07ryoushushosisutemu.xlsm",
]


def dump_workbook(path: Path) -> None:
    print("=" * 60)
    print(f"ファイル: {path.name}")
    print("=" * 60)
    wb = openpyxl.load_workbook(path, data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n--- シート: {sheet_name} (dims={ws.dimensions}) ---")
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and str(cell.value).strip() != "":
                    print(f"{cell.coordinate}: {cell.value!r}")


def main() -> None:
    for filename in TARGET_FILES:
        file_path = HAMAKURU_DIR / filename
        if not file_path.exists():
            print(f"[SKIP] ファイルが見つかりません: {filename}")
            continue
        dump_workbook(file_path)


if __name__ == "__main__":
    main()
